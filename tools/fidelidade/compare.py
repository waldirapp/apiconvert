"""Compara a conversão do xlsconvert com a do pandas (conversor de referência).

Uso (na máquina que tem o pandas do conversor de referência):

    python3 -I compare.py --rust-xlsx <saida_rust.xlsx> --entrada <arquivo.xls> [--valores]
    python3 -I compare.py --manifesto <pares.tsv> --resultado <saida.jsonl>

No modo manifesto, cada linha do TSV é ``entrada<TAB>saida_rust``; o pandas é
importado uma única vez, o JSON de cada arquivo vai para ``--resultado`` e só
o resumo agregado sai na saída padrão.

Gera a referência executando exatamente o que o conversor de referência executa:

    df = pd.read_excel(entrada)
    df.to_excel(ref, index=False, header=False)

e compara as duas saídas célula a célula com o openpyxl. Imprime um JSON
por arquivo. Sem ``--valores`` o relatório traz só posições e tipos das
diferenças, nunca o conteúdo: é o modo para arquivos de clientes.
"""

import argparse
import datetime as dt
import json
import math
import os
import sys
import tempfile


def _classe(v):
    if v is None or v == "":
        return "vazio", None
    if isinstance(v, bool):
        return "bool", v
    if isinstance(v, (int, float)):
        return "numero", float(v)
    if isinstance(v, dt.datetime):
        return "datahora", v
    if isinstance(v, dt.date):
        return "datahora", dt.datetime.combine(v, dt.time())
    if isinstance(v, dt.time):
        return "hora", v
    return "texto", str(v)


def _iguais(a, b):
    ca, va = _classe(a)
    cb, vb = _classe(b)
    if ca != cb:
        return False, ca, cb
    if ca == "numero":
        if math.isnan(va) and math.isnan(vb):
            return True, ca, cb
        return va == vb, ca, cb
    if ca == "datahora":
        # openpyxl reconstrói a data do serial; diferença de até 1 ms é
        # arredondamento do ponto flutuante, não de conteúdo.
        return abs((va - vb).total_seconds()) < 0.001, ca, cb
    return va == vb, ca, cb


def _grade(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=False)
    ws = wb.worksheets[0]
    cells = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and c.value != "":
                cells[(c.row, c.column)] = c.value
    return wb.sheetnames[0], len(wb.sheetnames), cells, ws


def comparar(entrada, rust_xlsx, mostrar_valores):
    import pandas as pd

    rel = {"arquivo": os.path.basename(entrada)}
    with tempfile.TemporaryDirectory() as tmp:
        ref = os.path.join(tmp, "ref.xlsx")
        try:
            df = pd.read_excel(entrada)
            df.to_excel(ref, index=False, header=False)
        except Exception as exc:  # o pandas recusou: o Rust também deve recusar
            rel["pandas"] = "erro"
            rel["pandas_erro"] = type(exc).__name__
            rel["rust"] = "erro" if not os.path.exists(rust_xlsx) else "convertido"
            rel["igual"] = rel["rust"] == "erro"
            return rel
        rel["pandas"] = "convertido"
        if not os.path.exists(rust_xlsx):
            rel["rust"] = "erro"
            rel["igual"] = False
            return rel
        rel["rust"] = "convertido"

        nome_ref, abas_ref, cel_ref, ws_ref = _grade(ref)
        nome_rust, abas_rust, cel_rust, ws_rust = _grade(rust_xlsx)

    difs = []
    if (nome_ref, abas_ref) != (nome_rust, abas_rust):
        difs.append({"tipo": "abas", "ref": [nome_ref, abas_ref], "rust": [nome_rust, abas_rust]})

    contagem = {}
    for pos in sorted(set(cel_ref) | set(cel_rust)):
        a = cel_ref.get(pos)
        b = cel_rust.get(pos)
        ok, ca, cb = _iguais(a, b)
        if ok:
            continue
        chave = f"{ca}->{cb}"
        contagem[chave] = contagem.get(chave, 0) + 1
        if len(difs) < 50:
            d = {"linha": pos[0], "coluna": pos[1], "ref": ca, "rust": cb}
            if mostrar_valores:
                d["valor_ref"] = repr(a)
                d["valor_rust"] = repr(b)
            difs.append(d)

    rel["celulas_ref"] = len(cel_ref)
    rel["celulas_rust"] = len(cel_rust)
    rel["diferencas_por_tipo"] = contagem
    rel["diferencas"] = difs
    rel["igual"] = not difs
    return rel


def _lote(manifesto, resultado, mostrar_valores):
    resumo = {"total": 0, "iguais": 0, "ambos_recusaram": 0, "falha_comparacao": 0,
              "diferencas_por_tipo": {}, "arquivos_diferentes": 0}
    with open(manifesto, encoding="utf-8") as m, open(resultado, "w", encoding="utf-8") as out:
        for linha in m:
            linha = linha.rstrip("\n")
            if not linha:
                continue
            entrada, rust_xlsx = linha.split("\t", 1)
            resumo["total"] += 1
            try:
                rel = comparar(entrada, rust_xlsx, mostrar_valores)
            except Exception as exc:  # arquivo que nem o openpyxl abre
                rel = {"arquivo": os.path.basename(entrada), "igual": False,
                       "falha_comparacao": type(exc).__name__}
                resumo["falha_comparacao"] += 1
            if rel.get("igual"):
                resumo["iguais"] += 1
                if rel.get("pandas") == "erro":
                    resumo["ambos_recusaram"] += 1
            elif "falha_comparacao" not in rel:
                resumo["arquivos_diferentes"] += 1
                for k, v in rel.get("diferencas_por_tipo", {}).items():
                    resumo["diferencas_por_tipo"][k] = resumo["diferencas_por_tipo"].get(k, 0) + v
                if rel.get("pandas") != rel.get("rust"):
                    chave = f"status:{rel.get('pandas')}->{rel.get('rust')}"
                    resumo["diferencas_por_tipo"][chave] = resumo["diferencas_por_tipo"].get(chave, 0) + 1
            out.write(json.dumps(rel, ensure_ascii=False, default=str) + "\n")
    print(json.dumps(resumo, ensure_ascii=False))
    return 0 if resumo["iguais"] == resumo["total"] else 1


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--entrada")
    p.add_argument("--rust-xlsx")
    p.add_argument("--manifesto")
    p.add_argument("--resultado")
    p.add_argument("--valores", action="store_true")
    a = p.parse_args()
    if a.manifesto:
        if not a.resultado:
            p.error("--manifesto exige --resultado")
        return _lote(a.manifesto, a.resultado, a.valores)
    if not (a.entrada and a.rust_xlsx):
        p.error("informe --entrada e --rust-xlsx, ou --manifesto")
    rel = comparar(a.entrada, a.rust_xlsx, a.valores)
    print(json.dumps(rel, ensure_ascii=False, default=str))
    return 0 if rel["igual"] else 1


if __name__ == "__main__":
    sys.exit(main())
